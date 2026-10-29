"""
PICANTE — review workbook for the manual reference supplement.

  1. build   python make_review_sheet.py build manual/match_log.csv manual/review.xlsx
             Turns every row that needs a human decision into an Excel workbook:
               Instructions  how to fill it in, your names, progress
               Review        one row per unique flagged reference (shared references appear once)
               Second check  the 10% verification sample + every hand match, for the second team member

  2. (fill in Review in Excel; save. Then the second checker fills in Second check; save.)
     Running build again on the same log is safe only before you start: it overwrites review.xlsx.

  3. merge   python make_review_sheet.py merge manual/match_log.csv manual/review.xlsx manual/match_log_reviewed.csv
             Writes your decisions into a copy of the match log, in the columns finalize reads.
             A DOI typed in the "OpenAlex ID or DOI" column is looked up and replaced by its W-ID
             (needs OPENALEX_API_KEY). Lists anything still missing and writes nothing until the
             Review tab is complete. Then run:
             python match_references.py finalize manual/match_log_reviewed.csv round1 manual

Needs: openpyxl (python3 -m pip install openpyxl), and match_references.py in the same folder.
"""
import csv, os, re, sys
from collections import OrderedDict
from datetime import date
from pathlib import Path
from urllib.parse import quote_plus

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from match_references import clean_title, norm_title  # same title cleaning as the matcher

DECISIONS = ["accept", "reject", "matched_by_hand", "not_in_openalex", "grey_literature"]
F = "Arial"
HEAD = PatternFill("solid", fgColor="1F3864"); HEADF = Font(name=F, bold=True, color="FFFFFF")
INPUT = PatternFill("solid", fgColor="FFF2CC")       # yellow = cells you fill in
DONE = PatternFill("solid", fgColor="E2EFDA")        # green = row decided
GROUPFILL = {"1": "DDEBF7", "2": "FCE4D6", "3": "EDEDED"}
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def classify(r):
    """Group 1: a candidate to confirm. Group 2: probably grey literature. Group 3: find by hand."""
    sim = float(r["similarity"]) if r["similarity"] else 0.0
    raw = r["raw"]
    has_cand = bool(r["openalex_id"])
    if r["route"] == "doi" and r["author_ok"] == "False":
        return "1", "DOI found it, but the first author differs. Check it is the same work."
    if has_cand and sim >= 0.80:
        why = []
        if sim < 0.90: why.append(f"titles {sim:.0%} similar (needs 90%)")
        if r["author_ok"] != "True": why.append("first author differs")
        return "1", "Close candidate: " + "; ".join(why) + ". Accept if it is the same work."
    webby = bool(re.search(r"https?://|www\.|Accessed|\[Online\]|Retrieved", raw)) and not re.search(r"10\.\d{4,9}/", raw)
    if webby or not r["year"] or not r["first_author"]:
        return "2", "Looks like a web page, report or other non-scholarly item. Mark grey_literature unless it is a paper or book."
    if has_cand:
        return "3", f"Best candidate is poor ({sim:.0%} similar). Search for the right record, or mark not_in_openalex."
    return "3", "No candidate found. Search for it, or mark not_in_openalex."


def key_of(r):
    t = norm_title(clean_title(r["title"]))
    return (t, r["year"][:4]) if len(t.split()) >= 3 else (norm_title(r["raw"])[:120], "")


def style_header(ws, headers, widths):
    ws.append(headers)
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i); c.fill = HEAD; c.font = HEADF; c.alignment = WRAP
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"


def build(log_path, out_path):
    rows = list(csv.DictReader(open(log_path, encoding="utf-8-sig")))
    flagged = [r for r in rows if r["needs_review"] == "True"]
    sample = [r for r in rows if r["in_verification_sample"] == "True"]

    groups = OrderedDict()
    for r in flagged:
        groups.setdefault(key_of(r), []).append(r)
    items = []
    for k, rs in groups.items():
        g, why = classify(rs[0])
        items.append((g, int(rs[0]["seed"]), int(rs[0]["ref_no"]), k, rs, why))
    items.sort(key=lambda x: (x[0], x[1], x[2]))

    wb = Workbook()
    # ------------------------------------------------------------ Instructions
    ins = wb.active; ins.title = "Instructions"
    ins.column_dimensions["A"].width = 34; ins.column_dimensions["B"].width = 95
    lines = [
        ("Manual reference supplement — review", None),
        ("", None),
        ("Your name (reviewer)", ""), ("Second checker's name", ""),
        ("", None),
        ("What to do", "On the Review tab, fill in the yellow cells: pick a Decision for every row. Rows turn green when decided."),
        ("", "If you choose matched_by_hand, paste the work's OpenAlex ID (W…) or its DOI in 'OpenAlex ID or DOI'. A DOI is looked up for you at merge."),
        ("", "Rows are sorted into three groups. Group 1 has a candidate to confirm (quick). Group 2 looks like grey literature (quick). Group 3 needs finding by hand."),
        ("", "A reference cited by several seeds appears once; 'Cited by' lists every seed and reference number, and your decision applies to all of them."),
        ("", "The second checker then fills in the Second check tab: the 10% verification sample, plus every row you matched by hand."),
        ("Decisions", "accept — the candidate shown is the right work"),
        ("", "reject — the candidate is wrong and you are not supplying another (the reference adds no link)"),
        ("", "matched_by_hand — you found the right record; give its W-ID or DOI"),
        ("", "not_in_openalex — a scholarly work OpenAlex does not have"),
        ("", "grey_literature — a web page, report, dataset, bill, talk or similar; it goes to the hand-searching list"),
        ("Example (not a real row)", "Decision = matched_by_hand · OpenAlex ID or DOI = 10.1145/3351095.3372873 · Notes = 'Endto-End' in PDF; FAT* 2020 paper"),
        ("", None),
        ("Progress", None),
        ("Review rows", "=COUNTA(Review!A2:A{n})"),
        ("Review rows decided", "=COUNTA(Review!J2:J{n})"),
        ("Review rows left", "=B19-B20"),
        ("Second-check rows to answer", "=COUNTIF('Second check'!A2:A{m},\"sample (auto match)\")+COUNTIF(Review!J2:J{n},\"matched_by_hand\")"),
        ("Second-check rows answered", "=COUNTA('Second check'!H2:H{m})"),
        ("", None),
        ("Source", f"Built {date.today().isoformat()} from {Path(log_path).name}: {len(flagged)} flagged rows collapsed to {len(items)} unique references; {len(sample)} rows in the verification sample."),
    ]
    for a, b in lines:
        ins.append([a, b])
    for row in ins.iter_rows():
        for c in row:
            c.font = Font(name=F); c.alignment = WRAP
    ins["A1"].font = Font(name=F, bold=True, size=14)
    for ref in ("B3", "B4"):
        ins[ref].fill = INPUT
    for ref in ("A6", "A11", "A16", "A18"):
        ins[ref].font = Font(name=F, bold=True)

    # ------------------------------------------------------------ Review
    rv = wb.create_sheet("Review")
    headers = ["Group", "Seed", "Ref #", "Cited by", "Reference as cited", "Why it needs you",
               "Candidate in OpenAlex", "Candidate link", "Find it (Google Scholar)",
               "Decision", "OpenAlex ID or DOI", "Notes", "key"]
    style_header(rv, headers, [7, 6, 6, 16, 60, 34, 40, 14, 14, 18, 22, 28, 4])
    for g, seed, ref, k, rs, why in items:
        r0 = rs[0]
        cited = "; ".join(f"seed {x['seed']} #{x['ref_no']}" for x in rs)
        cand = (f"{r0['matched_title']}  [{r0['openalex_id']}" + (f", {r0['similarity']}" if r0['similarity'] else "") + "]") if r0["openalex_id"] else "—"
        q = clean_title(r0["title"]) or r0["raw"][:200]
        rv.append([int(g), seed, ref, cited, r0["raw"], why, cand,
                   "open" if r0["openalex_id"] else "", "search", None, None, None,
                   "|".join(k)])
        n = rv.max_row
        if r0["openalex_id"]:
            rv.cell(n, 8).hyperlink = f"https://openalex.org/{r0['openalex_id']}"
        rv.cell(n, 9).hyperlink = "https://scholar.google.com/scholar?q=" + quote_plus(q + (" " + r0["first_author"] if r0["first_author"] else ""))
        for col in range(1, 14):
            c = rv.cell(n, col); c.font = Font(name=F, size=10); c.alignment = WRAP; c.border = THIN
        rv.cell(n, 1).fill = PatternFill("solid", fgColor=GROUPFILL[g])
        for col in (8, 9):
            rv.cell(n, col).font = Font(name=F, size=10, color="0563C1", underline="single")
        for col in (10, 11, 12):
            rv.cell(n, col).fill = INPUT
    last = max(rv.max_row, 2)
    rv.column_dimensions["M"].hidden = True
    dv = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Decision", error="Pick one of the listed decisions.")
    rv.add_data_validation(dv); dv.add(f"J2:J{last}")
    rv.conditional_formatting.add(f"A2:L{last}", FormulaRule(formula=[f'LEN($J2)>0'], fill=DONE))
    rv.auto_filter.ref = f"A1:L{last}"

    # ------------------------------------------------------------ Second check
    sc = wb.create_sheet("Second check")
    sh = ["Type", "Seed", "Ref #", "Reference as cited", "Record chosen", "Record link",
          "Reviewer's notes", "Same work? (Yes/No)", "If No: correct W-ID or DOI", "Notes", "key"]
    style_header(sc, sh, [16, 6, 6, 60, 44, 12, 26, 14, 22, 26, 4])
    for r in sorted(sample, key=lambda x: (int(x["seed"]), int(x["ref_no"]))):
        sc.append(["sample (auto match)", int(r["seed"]), int(r["ref_no"]), r["raw"],
                   f"{r['matched_title']}  [{r['openalex_id']}, {r['similarity']}]", "open", "", None, None, None,
                   f"sample|{r['seed']}|{r['ref_no']}"])
        sc.cell(sc.max_row, 6).hyperlink = f"https://openalex.org/{r['openalex_id']}"
    # hand matches are added from the Review tab by formula-free copy at merge time is not possible
    # in Excel without macros, so the tab lists every Review row; the checker answers only those
    # the Review tab marks matched_by_hand (the others turn grey).
    for i, (g, seed, ref, k, rs, why) in enumerate(items):
        rvrow = i + 2
        sc.append(["hand match?", seed, ref, rs[0]["raw"],
                   f"=IF(Review!J{rvrow}=\"matched_by_hand\",Review!K{rvrow},\"(not a hand match — skip)\")",
                   "", f"=IF(LEN(Review!L{rvrow})>0,Review!L{rvrow},\"\")", None, None, None,
                   f"hand|{'|'.join(k)}"])
    for row in sc.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=F, size=10); c.alignment = WRAP; c.border = THIN
        row[5].font = Font(name=F, size=10, color="0563C1", underline="single")
        for c in row[7:10]:
            c.fill = INPUT
    lastm = max(sc.max_row, 2)
    sc.column_dimensions["K"].hidden = True
    dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    sc.add_data_validation(dv2); dv2.add(f"H2:H{lastm}")
    sc.conditional_formatting.add(f"A2:J{lastm}", FormulaRule(formula=['LEN($H2)>0'], fill=DONE))
    sc.conditional_formatting.add(f"A2:J{lastm}", FormulaRule(formula=['ISNUMBER(SEARCH("skip",$E2))'],
                                                                fill=PatternFill("solid", fgColor="F2F2F2"),
                                                                font=Font(name=F, color="A6A6A6")))
    sc.auto_filter.ref = f"A1:J{lastm}"

    # progress formula ranges
    for row in ins.iter_rows(min_row=19, max_row=24, min_col=2, max_col=2):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                c.value = c.value.replace("{n}", str(last)).replace("{m}", str(lastm))
    wb.active = 1
    wb.save(out_path)
    print(f"{len(flagged)} flagged rows -> {len(items)} unique references "
          f"(group 1: {sum(1 for x in items if x[0]=='1')}, group 2: {sum(1 for x in items if x[0]=='2')}, "
          f"group 3: {sum(1 for x in items if x[0]=='3')}); {len(sample)} sample rows. Wrote {out_path}")


def resolve(ident, cache={}):
    """W-ID passes through; a DOI is looked up in OpenAlex (free single-work lookup)."""
    ident = (ident or "").strip()
    m = re.search(r"\b(W\d{4,})\b", ident)
    if m:
        return m.group(1), ""
    d = re.sub(r"^(https?://(dx\.)?doi\.org/|doi:\s*)", "", ident, flags=re.I).strip().rstrip(".")
    if not d.startswith("10."):
        return "", f"'{ident}' is neither a W-ID nor a DOI"
    if d in cache:
        return cache[d]
    import requests
    key = os.environ.get("OPENALEX_API_KEY")
    if not key:
        return "", "set OPENALEX_API_KEY to look up DOIs"
    r = requests.get(f"https://api.openalex.org/works/doi:{d}", params={"api_key": key, "select": "id"}, timeout=60)
    out = (r.json()["id"].rsplit("/", 1)[-1], "") if r.status_code == 200 else ("", f"DOI {d} not found in OpenAlex")
    cache[d] = out
    return out


def merge(log_path, xlsx_path, out_path):
    rows = list(csv.DictReader(open(log_path, encoding="utf-8-sig")))
    wb = load_workbook(xlsx_path, data_only=True)
    ins, rv, sc = wb["Instructions"], wb["Review"], wb["Second check"]
    reviewer = (ins["B3"].value or "").strip()
    checker = (ins["B4"].value or "").strip()
    problems = []
    if not reviewer:
        problems.append("Instructions!B3: reviewer's name is empty")

    decided = {}
    for row in rv.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        seed, ref, dec, ident, notes, key = row[1], row[2], (row[9] or "").strip(), row[10], row[11], row[12]
        if dec not in DECISIONS:
            problems.append(f"Review seed {seed} #{ref}: no decision" if not dec else f"Review seed {seed} #{ref}: unknown decision '{dec}'")
            continue
        wid = ""
        if dec == "matched_by_hand":
            wid, err = resolve(str(ident or ""))
            if err:
                problems.append(f"Review seed {seed} #{ref}: {err}")
                continue
        if dec == "accept" and "—" == (row[6] or "—"):
            problems.append(f"Review seed {seed} #{ref}: 'accept' but there is no candidate")
            continue
        decided[key] = (dec, wid, notes or "")

    checks = {}
    n_sample = n_sample_yes = n_hand = n_hand_yes = 0
    for row in sc.iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        kind, seed, ref, chosen, ans, fix, key = row[0], row[1], row[2], row[4], (row[7] or "").strip(), row[8], row[10]
        is_hand = key.startswith("hand|")
        if is_hand and decided.get(key[5:], ("",))[0] != "matched_by_hand":
            continue                                   # not a hand match: nothing to check
        if ans not in ("Yes", "No"):
            problems.append(f"Second check seed {seed} #{ref}: not answered")
            continue
        if is_hand:
            n_hand += 1; n_hand_yes += ans == "Yes"
        else:
            n_sample += 1; n_sample_yes += ans == "Yes"
        wid = ""
        if ans == "No" and fix:
            wid, err = resolve(str(fix))
            if err:
                problems.append(f"Second check seed {seed} #{ref}: {err}"); continue
        checks[key] = (ans, wid)
    if checks and not checker:
        problems.append("Instructions!B4: second checker's name is empty")

    if problems:
        print("Not merged. Fix these in the workbook and run merge again:")
        for p in problems[:60]:
            print("  -", p)
        if len(problems) > 60:
            print(f"  … and {len(problems) - 60} more")
        sys.exit(1)

    today = date.today().isoformat()
    by_key = {}
    for r in rows:
        if r["needs_review"] == "True":
            by_key.setdefault("|".join(key_of(r)), []).append(r)
    for key, (dec, wid, notes) in decided.items():
        for r in by_key.get(key, []):
            r["final_decision"], r["manual_openalex_id"] = dec, wid
            r["reviewer"], r["review_date"], r["notes"] = reviewer, today, notes
            ck = checks.get("hand|" + key)
            if ck:
                r["notes"] = (r["notes"] + "; " if r["notes"] else "") + f"second check ({checker}): {ck[0]}"
                if ck[0] == "No":
                    r["final_decision"], r["manual_openalex_id"] = ("matched_by_hand", ck[1]) if ck[1] else ("reject", "")
                    r["notes"] += "; corrected at second check"
    for r in rows:
        ck = checks.get(f"sample|{r['seed']}|{r['ref_no']}")
        if ck:
            r["reviewer"], r["review_date"] = checker, today
            r["final_decision"] = "accept" if ck[0] == "Yes" else ("matched_by_hand" if ck[1] else "reject")
            r["manual_openalex_id"] = ck[1] if ck[0] == "No" else ""
            r["notes"] = f"verification sample: {ck[0]}"

    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, rows[0].keys()); w.writeheader(); w.writerows(rows)
    print(f"Merged {len(decided)} reviewed references into {out_path}.")
    if n_sample:
        print(f"Verification sample: {n_sample_yes}/{n_sample} automatic title matches confirmed ({n_sample_yes / n_sample:.0%}).")
    if n_hand:
        print(f"Hand matches confirmed at second check: {n_hand_yes}/{n_hand}.")
    print(f"Next: python match_references.py finalize {out_path} round1 manual")


if __name__ == "__main__":
    if len(sys.argv) >= 4 and sys.argv[1] == "build":
        build(sys.argv[2], sys.argv[3])
    elif len(sys.argv) >= 5 and sys.argv[1] == "merge":
        merge(sys.argv[2], sys.argv[3], sys.argv[4])
    else:
        sys.exit(__doc__)
