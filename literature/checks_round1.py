"""
PICANTE literature strand -- checks on round-1 expansion output.

Run after expansion_round1.py. Reads its output folder plus the seed and
validation files, makes a handful of free single-work lookups, and writes
checks_report.md for the record.

Usage:
    export OPENALEX_API_KEY=...
    python checks_round1.py TPG_seed_set_v3.xlsx TPG_search_validation_set.xlsx round1/
"""
import csv, json, os, re, sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from difflib import SequenceMatcher
from pathlib import Path
import dotenv
dotenv.load_dotenv()

import openpyxl
import requests

KEY = os.environ.get("OPENALEX_API_KEY") or sys.exit("Set OPENALEX_API_KEY first.")
seed_file, val_file, out = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
log = open(out / "access_log.jsonl", "a")


def lookup(doi):
    """Single-work DOI lookup (free). Returns the record or None if not indexed."""
    r = requests.get(f"https://api.openalex.org/works/doi:{doi}",
                     params={"api_key": KEY, "select": "id,doi,display_name,locations,cited_by_count"},
                     timeout=60)
    log.write(json.dumps({"utc": datetime.now(timezone.utc).isoformat(),
                          "path": f"/works/doi:{doi}", "status": r.status_code}) + "\n")
    if r.status_code == 404:
        return None
    r.raise_for_status()
    return r.json()


def norm(t):
    return re.sub(r"[^a-z0-9 ]", "", (t or "").lower())


def sheet(path, name):
    ws = openpyxl.load_workbook(path, data_only=True)[name]
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


seeds = {int(r["Seed #"]): r for r in sheet(seed_file, "Seed set (23)")}
vals = sheet(val_file, "Validation set (8)")
resolved = {int(r["seed"]): r for r in csv.DictReader(open(out / "seeds_resolved.csv"))}
yields = {int(r["seed"]): r for r in csv.DictReader(open(out / "yields.csv"))}
pool = {r["openalex_id"]: r for r in csv.DictReader(open(out / "pool_prededup.csv"))}
reached = defaultdict(list)          # work -> [(seed, direction)]
for e in csv.DictReader(open(out / "edges.csv")):
    other = e["cited_id"] if e["direction"] == "backward" else e["citing_id"]
    reached[other].append((e["seed"], e["direction"]))

R = [f"# Round-1 checks\n\nRun {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC.\n"]

# ---- Check 1: every seed resolved to the right work --------------------------
R.append("## 1. Seed resolution\n\n| Seed | OpenAlex ID | Title match | Flag |\n|---|---|---|---|")
pool_titles = {k: v["title"] for k, v in pool.items()}
for n, s in sorted(seeds.items()):
    if n not in resolved:
        R.append(f"| {n} | — | — | not resolved (expected only for seed 20) |"); continue
    wid = resolved[n]["openalex_id"]
    sim = SequenceMatcher(None, norm(s["Title"]), norm(pool_titles.get(wid))).ratio()
    R.append(f"| {n} | {wid} | {sim:.2f} | {'CHECK BY HAND' if sim < 0.9 else ''} |")

# ---- Check 2: per-seed yields --------------------------------------------------
R.append("\n## 2. Per-seed yields\n\nForward should equal cited_by_count (a gap means "
         "paging stopped early or the index changed mid-run). Zero backward on an indexed "
         "seed means OpenAlex holds no reference list for it: record as materially "
         "incomplete (prereg 1.6).\n\n| Seed | Backward | Forward | cited_by_count | Viability (seed file) | Flag |\n|---|---|---|---|---|---|")
for n, y in sorted(yields.items()):
    b, f, c = int(y["backward"]), int(y["forward"]), int(resolved[n]["cited_by_count"])
    viab = str(seeds[n]["Forward-chase viability"] or "")[:20]
    flags = []
    if b == 0: flags.append("no references indexed")
    if f != c: flags.append(f"forward ≠ cited_by_count ({c - f:+d})")
    if f < 5 and viab.startswith("Expected normal"): flags.append("low forward for a 'normal' seed")
    R.append(f"| {n} | {b} | {f} | {c} | {viab} | {'; '.join(flags)} |")

# ---- Check 3: preprint / published merges --------------------------------------
R.append("\n## 3. Preprint and published records\n")
rec = lookup("10.48550/arXiv.2505.22327")
s22 = resolved.get(22, {}).get("openalex_id")
if rec is None:
    R.append("Seed 22 preprint: arXiv DOI not indexed separately.")
else:
    pid = rec["id"].rsplit("/", 1)[-1]
    R.append(f"Seed 22: published record {s22}, preprint record {pid} -> "
             + ("**MERGED** (same work)." if pid == s22 else
                f"**NOT MERGED**. Preprint has {rec['cited_by_count']} citing works not in "
                "seed 22's forward set; report its forward set as incomplete (prereg 2.5)."))
for n in (6, 13, 23):
    if n in resolved:
        R.append(f"\nSeed {n} (arXiv): cited_by_count {resolved[n]['cited_by_count']}.")

# ---- Check 4: validation set against the PRE-dedup pool -----------------------
R.append("\n## 4. Search validation (pre-deduplication pool)\n\n"
         "| V | Arm | Counts | OpenAlex ID | In pool | Reached via | Status |\n|---|---|---|---|---|---|---|")
c_count = 0
for v in vals:
    vid, arm, counts = v["V #"], v["Arm tested"], str(v["Counts toward expansion criterion"])
    ident = str(v["DOI / stable ID"])
    if not ident.startswith("10."):
        R.append(f"| {vid} | {arm} | {counts[:3]} | — | — | — | hand-searching arm: check at that stage |"); continue
    rec = lookup(ident)
    if rec is None:
        R.append(f"| {vid} | {arm} | {counts[:3]} | — | — | — | (a) absent from index |"); continue
    wid = rec["id"].rsplit("/", 1)[-1]
    hit = wid in pool
    via = ", ".join(f"seed {s} {d}" for s, d in reached.get(wid, []))
    status = "recovered" if hit else "not yet reached (would be (c) if still absent at saturation)"
    if not hit and counts.startswith("Yes"): c_count += 1
    R.append(f"| {vid} | {arm} | {counts[:3]} | {wid} | {hit} | {via} | {status} |")
R.append(f"\nExpansion-counting works (V1–V5) not reached in round 1: **{c_count}**. "
         "Threshold for adding a targeted query: 2 unreachable at saturation.")

# ---- Check 5: pool quality ----------------------------------------------------
n = len(pool)
types = Counter(r["type"] for r in pool.values())
langs = Counter(r["language"] for r in pool.values())
no_abs = sum(1 for r in pool.values() if not r["abstract"])
no_doi = sum(1 for r in pool.values() if not r["doi"])
retr = [k for k, r in pool.items() if r["is_retracted"] == "True"]
titles = Counter(norm(r["title"]) for r in pool.values() if r["title"])
dup_titles = {t: c for t, c in titles.items() if c > 1}
R.append(f"\n## 5. Pool quality\n\n- Records (pre-dedup, seeds included): {n}\n"
         f"- Missing abstract: {no_abs} ({no_abs / n:.0%}) — these need full text at Stage 1\n"
         f"- Missing DOI: {no_doi}\n"
         f"- Flagged retracted by OpenAlex: {len(retr)} {retr[:10]}\n"
         f"- Normalized titles appearing more than once: {len(dup_titles)} (candidate duplicates for the human check)\n"
         f"- Types: {dict(types.most_common())}\n- Languages: {dict(langs.most_common(8))}\n")

(out / "checks_report.md").write_text("\n".join(R))
print(f"Wrote {out / 'checks_report.md'}")
